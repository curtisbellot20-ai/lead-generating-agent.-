import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from lead_gen import config

HEADERS = [
    "Business Name",
    "Owner Name",
    "Decision Maker",
    "DM Title",
    "Phone",
    "Email",
    "Address",
    "Website",
    "Instagram",
    "Facebook",
    "TikTok",
    "LinkedIn",
    "Category",
    "Rating",
    "Reviews",
    "Years in Biz",
    "Description",
    "Industry Tag",
    "Registry Agent",
    "Registry Member",
    "Registry Status",
    "Source",
    "Outreach Note",
]

COL_WIDTHS = [
    35, 22, 22, 18,          # Business Name, Owner, DM, DM Title
    18, 34,                  # Phone, Email
    45, 35,                  # Address, Website
    30, 30, 26, 30,          # Instagram, Facebook, TikTok, LinkedIn
    18, 8, 8, 10,            # Category, Rating, Reviews, Years
    55, 18,                  # Description, Industry Tag
    22, 22, 14,              # Registry Agent, Member, Status
    16, 65,                  # Source, Outreach Note
]

# Column indices (1-based) for coloured text
_COL = {h: i for i, h in enumerate(HEADERS, 1)}
_LINK_COLORS = {
    _COL["Email"]:      "0563C1",
    _COL["Instagram"]:  "C13584",
    _COL["Facebook"]:   "1877F2",
    _COL["TikTok"]:     "010101",
    _COL["LinkedIn"]:   "0A66C2",
}
# Registry columns get a light-blue background to show they are verified
_REGISTRY_COLS = {_COL["Registry Agent"], _COL["Registry Member"], _COL["Registry Status"]}
_REG_FILL = PatternFill(start_color="EBF5FB", end_color="EBF5FB", fill_type="solid")

# Owner Name column confidence colours
_CONF_FONT = {
    "HIGH":   Font(color="1A7A1A", bold=True),   # green
    "MEDIUM": Font(color="7D4B00"),              # amber
    "LOW":    Font(color="999999"),              # grey
}


def _border():
    s = Side(style="thin", color="DDDDDD")
    return Border(left=s, right=s, top=s, bottom=s)


def export_to_excel(leads: list[dict], query: str, location: str) -> str:
    os.makedirs(config.OUTPUT_DIR, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    slug      = query.lower().replace(" ", "_")
    filename  = f"{config.OUTPUT_DIR}/leads_{slug}_{timestamp}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    dark_blue = "0F3460"
    mid_blue  = "1A1A2E"
    alt_gray  = "F8F9FA"
    border    = _border()
    last_col  = get_column_letter(len(HEADERS))

    # ── Title row
    ws.merge_cells(f"A1:{last_col}1")
    tc = ws["A1"]
    tc.value     = f"Lead Generation Report — {query.title()} in {location}"
    tc.font      = Font(bold=True, color="FFFFFF", size=13)
    tc.fill      = PatternFill(start_color=dark_blue, end_color=dark_blue, fill_type="solid")
    tc.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    # ── Meta row
    emails_found   = sum(1 for l in leads if l.get("email"))
    social_found   = sum(1 for l in leads if l.get("instagram") or l.get("facebook") or l.get("tiktok"))
    registry_found = sum(1 for l in leads if l.get("registry_agent") or l.get("registry_member"))
    ws.merge_cells(f"A2:{last_col}2")
    mc = ws["A2"]
    mc.value = (
        f"Generated: {datetime.now().strftime('%B %d, %Y %I:%M %p')}  "
        f"|  Leads: {len(leads)}  "
        f"|  Emails: {emails_found}  "
        f"|  Social: {social_found}  "
        f"|  Registry matches: {registry_found}"
    )
    mc.font      = Font(italic=True, color="666666", size=10)
    mc.alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 20

    # ── Header row
    hdr_fill  = PatternFill(start_color=mid_blue, end_color=mid_blue, fill_type="solid")
    hdr_font  = Font(bold=True, color="FFFFFF", size=11)
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, (hdr, width) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
        cell           = ws.cell(row=3, column=col, value=hdr)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = hdr_align
        cell.border    = border
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[3].height = 30

    # ── Data rows
    data_align = Alignment(vertical="center", wrap_text=True)
    owner_col  = _COL["Owner Name"]

    for row_idx, lead in enumerate(leads, 4):
        row_fill = (
            PatternFill(start_color=alt_gray, end_color=alt_gray, fill_type="solid")
            if row_idx % 2 == 0 else None
        )
        values = [
            lead.get("name", ""),
            lead.get("owner_name", ""),
            lead.get("decision_maker", ""),
            lead.get("decision_maker_title", ""),
            lead.get("phone", ""),
            lead.get("email", ""),
            lead.get("address", ""),
            lead.get("website", ""),
            lead.get("instagram", ""),
            lead.get("facebook", ""),
            lead.get("tiktok", ""),
            lead.get("linkedin", ""),
            lead.get("category", ""),
            lead.get("rating", ""),
            lead.get("reviews", ""),
            lead.get("years_in_business", ""),
            lead.get("description", ""),
            lead.get("industry_tag", ""),
            lead.get("registry_agent", ""),
            lead.get("registry_member", ""),
            lead.get("registry_status", ""),
            lead.get("source", ""),
            lead.get("outreach_note", ""),
        ]
        for col, value in enumerate(values, 1):
            cell           = ws.cell(row=row_idx, column=col, value=value)
            cell.alignment = data_align
            cell.border    = border
            if row_fill:
                cell.fill = row_fill

            # Social / email link colour
            if col in _LINK_COLORS and value:
                cell.font = Font(color=_LINK_COLORS[col], underline="single")
            # Owner confidence colour
            elif col == owner_col and value:
                conf = lead.get("owner_confidence", "")
                if conf in _CONF_FONT:
                    cell.font = _CONF_FONT[conf]
            # Registry verified background
            if col in _REGISTRY_COLS and value:
                cell.fill = _REG_FILL

        ws.row_dimensions[row_idx].height = 50

    ws.freeze_panes = "A4"

    # ── Summary sheet
    ws2 = wb.create_sheet("Summary")
    ws2["A1"].value = f"Summary — {query.title()} in {location}"
    ws2["A1"].font  = Font(bold=True, size=14)
    ws2["A2"].value = f"{len(leads)} leads  |  {emails_found} emails  |  {social_found} social  |  {registry_found} registry matches"
    ws2["A2"].font  = Font(italic=True, color="666666")

    sources: dict[str, int] = {}
    industries: dict[str, int] = {}
    for lead in leads:
        src = lead.get("source", "Unknown")
        sources[src] = sources.get(src, 0) + 1
        ind = lead.get("industry_tag") or "Unknown"
        industries[ind] = industries.get(ind, 0) + 1

    row = 4
    ws2.cell(row=row, column=1, value="Leads by Source").font = Font(bold=True)
    row += 1
    for src, cnt in sorted(sources.items(), key=lambda x: -x[1]):
        ws2.cell(row=row, column=1, value=src)
        ws2.cell(row=row, column=2, value=cnt)
        row += 1
    row += 1
    ws2.cell(row=row, column=1, value="Leads by Industry").font = Font(bold=True)
    row += 1
    for ind, cnt in sorted(industries.items(), key=lambda x: -x[1]):
        ws2.cell(row=row, column=1, value=ind)
        ws2.cell(row=row, column=2, value=cnt)
        row += 1

    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 15

    wb.save(filename)
    print(f"  [Exporter] Saved: {filename}")
    return filename
